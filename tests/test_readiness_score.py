# SPDX-License-Identifier: Apache-2.0
"""One readiness score, computed in code from the per-criterion verdicts and the gate's weights, the
same in the HIL task, the ICD and the Milestone package; dates filled; DoDAF views named correctly;
Word output without raw Markdown (JOB-20261008-DFA941: 40/100 "corrected" to 70/100 where the
weights give 60, "[Current Date]", OV-1 titled "Capability View", ** and <br> in Word tables)."""

import io
from types import SimpleNamespace

from docx import Document

from k9_dow.gates.gate_registry import DAS_GATES
from k9_dow.gates.readiness import (align_score, apply_computed_score, criterion_verdicts, fill_date,
                                    readiness_score)
from k9_dow.gates.review_summary import summarize_readiness
from k9_dow.utils.dodaf_views import label_view, relabel_views
from k9_dow.utils.icd_composer import nest_section, with_computed_score

JROC = DAS_GATES["JROC-VALIDATION"].entry_criteria

# The DFA941 scorer output, trimmed: table style, a model-written score, its own recalculation.
DFA941 = """# JROC-VALIDATION Gate Readiness Assessment
**Assessment Date:** [Current Date]
**Overall Readiness Score:** 40/100
| # | Criterion | Score | Rationale & Evidence Reference |
| :--- | :--- | :--- | :--- |
| 1 | **Capability need statement complete** | **MET** | **Evidence:** `model_elements`<br>complete |
| 2 | **Requirements traceability coverage >= 90%** | **MET** | coverage 92% |
| 3 | **DoDAF views generated and consistency-checked** | **PARTIALLY_MET** | only OV-1 |
| 4 | **No critical invariant violations** | **NOT_MET** | two critical incoherences |
| 5 | **Evidence package assembled** | **MET** | assembled |
## Weighted Overall Score Calculation
*   **Criterion 1 (MET):** 20 points (Weight: 20%)
**Total Score:** 20 + 20 + 10 + 0 + 20 = **70/100**
## Recommendation
*Correction:* The initial score of 40/100 was a conservative estimate.
Do not proceed.
"""

# List style (the Milestone package's readiness assessment).
LIST_STYLE = """Criterion 1: JROC validation approved
• Score: MET
Criterion 2: Pathway recommendation prepared
• Score: NOT_MET
Criterion 3: Funding line identified
• Score: NOT_MET
Criterion 4: JCIDS review package available (ICD, readiness assessment, artifact manifest)
• Score: MET
"""


def test_score_is_computed_from_verdicts_and_weights():
    out = apply_computed_score(DFA941, "JROC-VALIDATION", JROC, "2026-10-08")
    assert out["score"] == 60                                   # 20 + 20 + 10 + 0 + 10
    text = out["output"]
    assert "**Overall Readiness Score:** 60/100" in text
    assert "40/100" not in text and "70/100" not in text and "Correction" not in text
    assert "Score Calculation" not in text and "[Current Date]" not in text
    assert text.startswith("# JROC-VALIDATION Gate Readiness Assessment\n")   # block goes after the title
    assert "| 4 | No critical invariant violations | Not met | 30% | 0 |" in text


def test_list_style_verdicts_and_equal_weights():
    criteria = DAS_GATES["PATHWAY-MILESTONE"].entry_criteria
    assert criterion_verdicts(LIST_STYLE, criteria) == ["MET", "NOT_MET", "NOT_MET", "MET"]
    assert apply_computed_score(LIST_STYLE, "PATHWAY-MILESTONE", criteria)["score"] == 50


def test_missing_verdict_is_never_guessed():
    out = apply_computed_score("Criterion 1: Capability need statement complete — MET", "JROC-VALIDATION", JROC)
    assert out["score"] is None and "not computed" in out["output"]
    assert readiness_score([], []) is None


def test_downstream_text_states_the_computed_score():
    gap = "**Overall Readiness Score:** 40/100\n**Readiness Score:** 60/100 (Conservative) / 70/100 (Weighted)\n" \
          "*Correction:* 70 is right.\nDate: [Current Date]"
    out = align_score(gap, 60, "2026-10-08")
    assert out.count("60/100") == 2 and "40/100" not in out and "70/100" not in out
    assert "Correction" not in out and "2026-10-08" in out
    assert fill_date("on [Date]", "2026-01-02") == "on 2026-01-02"


def test_hil_summary_shows_the_computed_score():
    assert summarize_readiness(DFA941, JROC, 60)["Readiness"] == "60 / 100"
    assert summarize_readiness(DFA941, JROC)["Readiness"] == "40 / 100"      # without it: what the text says


def test_scorer_agent_returns_the_computed_score(monkeypatch):
    from k9_dow.agents.src import readiness_scorer_agent as mod
    seen = {}
    monkeypatch.setattr(mod, "llm_invoke", lambda cfg, req: seen.setdefault("p", req.prompt) and SimpleNamespace(output=DFA941))
    out = mod.ReadinessScorerAgent(config={}).execute({"gate_id": "JROC-VALIDATION", "gate_criteria": JROC})
    assert out["score"] == 60 and out["verdicts"]["No critical invariant violations"] == "NOT_MET"
    assert "Do NOT compute or state an overall score" in seen["p"] and "Assessment date:" in seen["p"]


def test_gap_reporter_restates_the_computed_score(monkeypatch):
    from k9_dow.agents.src import gap_reporter_agent as mod
    monkeypatch.setattr(mod, "llm_invoke", lambda cfg, req: SimpleNamespace(output="**Overall Readiness Score:** 40/100"))
    out = mod.GapReporterAgent(config={}).execute({"gate_criteria": JROC, "readiness_score": {"output": "x", "score": 60}})
    assert out["output"] == "**Overall Readiness Score:** 60/100"


def test_stored_jobs_get_the_computed_score_when_composed():
    stored = {"job_id": "JOB-20261008-DFA941", "gate_readiness": {
        "criteria": {"criteria": JROC, "output": "Loaded 5"},
        "readiness_score": {"output": DFA941},
        "gap_report": {"output": "**Overall Readiness Score:** 40/100"}},
        "review_package": {"review_package": {"output": "Readiness Score: 70/100"}}}
    out = with_computed_score(stored, "JROC-VALIDATION")
    assert out["gate_readiness"]["readiness_score"]["score"] == 60
    assert "60/100" in out["gate_readiness"]["gap_report"]["output"]
    assert out["review_package"]["review_package"]["output"] == "Readiness Score: 60/100"
    assert "2026-10-08" in out["gate_readiness"]["readiness_score"]["output"]


def test_views_carry_their_dodaf_names():
    view = "# DoDAF 2.0 View: OV-1 (Capability View)\n\n**View ID:** OV-1\n**View Name:** Capability View\n"
    assert label_view(view, "OV-1") == ("# DoDAF 2.0 View: OV-1 (High-Level Operational Concept Graphic)\n\n"
                                        "**View ID:** OV-1\n**View Name:** High-Level Operational Concept Graphic\n")
    assert relabel_views("Only `OV-1` (Capability View); the CV (Capability View)") == \
        "Only `OV-1` (High-Level Operational Concept Graphic); the CV (Capability View)"


def test_agent_title_is_not_repeated_under_the_section_heading():
    assert nest_section("# Readiness Assessment\n\n## Blockers\n- x", 3) == "#### Blockers\n- x"


def test_word_has_no_raw_markdown_and_list_numbers_restart():
    from k9_dow.reporting.docx.md_document import markdown_to_docx
    md = ("## A\n1. one\n2. two\n## B\n1. again\n\n| Criterion | Verdict |\n| --- | --- |\n"
          "| **Funding** | **NOT_MET**<br>`funding_line` missing |\n\n- **bold** and `code`\n    * nested")
    doc = Document(io.BytesIO(markdown_to_docx(md)))
    texts = [p.text for p in doc.paragraphs]
    assert "1.\tagain" in texts                                  # numbering restarts per list
    cell = doc.tables[0].cell(1, 1)
    assert cell.text == "NOT_MET\nfunding_line missing"         # <br> is a line break, ** and ` are gone
    assert cell.paragraphs[0].runs[0].bold
    assert not any(m in t for t in texts for m in ("**", "`", "<br>"))
    assert any(p.style.name == "List Bullet 2" for p in doc.paragraphs)
